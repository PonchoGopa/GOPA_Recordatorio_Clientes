import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def create_fiscal_reminder_template():
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # 1. PESTAÑA: CONFIGURACIÓN
    # -------------------------------------------------------------
    ws_config = wb.active
    ws_config.title = "Configuracion"
    ws_config.views.sheetView[0].showGridLines = True
    
    # Estilos de colores
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid") # Navy institucional
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=14, bold=True, color="1B365D")
    subtitle_font = Font(name="Segoe UI", size=9, italic=True, color="555555")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    normal_font = Font(name="Segoe UI", size=10)
    note_font = Font(name="Segoe UI", size=9, italic=True, color="666666")
    
    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )
    
    # Título Configuracion
    ws_config["A1"] = "CONFIGURACIÓN DEL DESPACHO CONTABLE"
    ws_config["A1"].font = title_font
    ws_config["A2"] = "Edite los valores en la columna B. n8n leerá automáticamente estos parámetros para personalizar correos y WhatsApp."
    ws_config["A2"].font = subtitle_font
    
    config_headers = ["Parámetro de Configuración", "Valor Configurado", "Descripción / Instrucciones"]
    for col_idx, text in enumerate(config_headers, start=1):
        cell = ws_config.cell(row=4, column=col_idx, value=text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    config_data = [
        ("Nombre del Despacho", "GOPA Asesores Fiscales", "Nombre comercial que aparecerá en los correos y mensajes de WhatsApp"),
        ("Teléfono de Contacto", "+52 664 123 4567", "Teléfono de atención al cliente que se muestra en la firma del mensaje"),
        ("Correo para Reportes", "contacto@gopa.mx", "Correo del contador donde se recibe el reporte semanal de vencimientos"),
        ("Días de Anticipación", "7, 3, 1, 0", "Días antes del vencimiento en que se envían los recordatorios (separados por coma)"),
        ("Enviar por Correo", "SI", "SI para enviar recordatorios por Correo (Gmail u Outlook), NO para desactivar"),
        ("Enviar por WhatsApp", "SI", "SI para enviar recordatorios por WhatsApp, NO para desactivar"),
        ("Hora de Envío Diario", "08:00", "Hora en la que n8n ejecuta la revisión diaria de recordatorios")
    ]
    
    for row_idx, (param, val, desc) in enumerate(config_data, start=5):
        c1 = ws_config.cell(row=row_idx, column=1, value=param)
        c2 = ws_config.cell(row=row_idx, column=2, value=val)
        c3 = ws_config.cell(row=row_idx, column=3, value=desc)
        
        c1.font = bold_font
        c2.font = normal_font
        c3.font = note_font
        
        c1.border = thin_border
        c2.border = thin_border
        c3.border = thin_border
        
        c2.alignment = Alignment(horizontal="center")
        if param in ["Enviar por Correo", "Enviar por WhatsApp"]:
            c2.fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    
    # Validación SI/NO en B9 y B10
    dv_sino = DataValidation(type="list", formula1='"SI,NO"', allow_blank=False)
    ws_config.add_data_validation(dv_sino)
    dv_sino.add("B9")
    dv_sino.add("B10")

    # -------------------------------------------------------------
    # 2. PESTAÑA: OBLIGACIONES_CLIENTES
    # -------------------------------------------------------------
    ws_clients = wb.create_sheet(title="Obligaciones_Clientes")
    ws_clients.views.sheetView[0].showGridLines = True
    
    ws_clients["A1"] = "CONTROL DE VENCIMIENTOS Y RECORDATORIOS FISCALES"
    ws_clients["A1"].font = title_font
    ws_clients["A2"] = "Agregue una fila por cada obligación de cada cliente. La columna 'Estado Vencimiento' se calcula automáticamente."
    ws_clients["A2"].font = subtitle_font
    
    columns = [
        ("ID", "Identificador único"),
        ("Cliente_Nombre", "Nombre de la empresa o cliente"),
        ("RFC", "RFC del contribuyente"),
        ("Cliente_Email", "Correo de contacto del cliente"),
        ("Telefono_WhatsApp", "Teléfono internacional (ej. +52 664 123 4567)"),
        ("Obligacion", "Tipo de declaración u obligación"),
        ("Periodo", "Mes o periodo a presentar"),
        ("Fecha_Vencimiento", "Fecha límite (YYYY-MM-DD)"),
        ("Estado_Vencimiento", "Cálculo automático de estatus"),
        ("Ultima_Fecha_Envio", "Última fecha en que n8n envió alerta"),
        ("Estatus_Envio", "Resultado de la última notificación"),
        ("Notas", "Observaciones del despacho")
    ]
    
    client_header_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
    calc_header_fill = PatternFill(start_color="4A6572", end_color="4A6572", fill_type="solid")
    
    for col_idx, (col_name, _) in enumerate(columns, start=1):
        cell = ws_clients.cell(row=4, column=col_idx, value=col_name)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if col_name in ["Estado_Vencimiento", "Ultima_Fecha_Envio", "Estatus_Envio"]:
            cell.fill = calc_header_fill
        else:
            cell.fill = client_header_fill
    
    sample_clients = [
        (1, "Comercializadora del Norte S.A. de C.V.", "CNO180415AA1", "facturacion@comercializadoranorte.com", "+52 664 123 4567", "Declaración Mensual IVA", "Agosto 2026", "2026-09-17", "Pendiente", "Cliente preferente, enviar confirmación."),
        (2, "Comercializadora del Norte S.A. de C.V.", "CNO180415AA1", "facturacion@comercializadoranorte.com", "+52 664 123 4567", "Pago Provisional ISR", "Agosto 2026", "2026-09-17", "Pendiente", "Solicitar balanza previa."),
        (3, "Dr. Roberto Silva Morales", "SIMR820914KL8", "roberto.silva@clinicadental.mx", "+52 55 9876 5432", "Declaración Informativa DIOT", "Agosto 2026", "2026-09-20", "Pendiente", "Faltan facturas de gastos."),
        (4, "Innovaciones Tecnológicas Nexo S.A.S.", "ITN210202H11", "contacto@nexotec.io", "+52 33 4455 6677", "Retenciones de ISR / IVA", "Agosto 2026", "2026-09-17", "Pendiente", "Pagar antes de las 14:00 hrs."),
        (5, "Servicios Gastronómicos La Costa S. de R.L.", "SGC150830TR4", "admon@lacostarestaurante.com", "+52 81 2233 4455", "Pago Provisional ISR", "Agosto 2026", "2026-09-17", "Pendiente", "Revisar retenciones RESICO.")
    ]
    
    for row_idx, item in enumerate(sample_clients, start=5):
        c_id, c_nom, c_rfc, c_email, c_tel, c_ob, c_per, c_fec, c_est, c_not = item
        
        ws_clients.cell(row=row_idx, column=1, value=c_id).alignment = Alignment(horizontal="center")
        ws_clients.cell(row=row_idx, column=2, value=c_nom)
        ws_clients.cell(row=row_idx, column=3, value=c_rfc).alignment = Alignment(horizontal="center")
        ws_clients.cell(row=row_idx, column=4, value=c_email)
        ws_clients.cell(row=row_idx, column=5, value=c_tel).alignment = Alignment(horizontal="center")
        ws_clients.cell(row=row_idx, column=6, value=c_ob)
        ws_clients.cell(row=row_idx, column=7, value=c_per).alignment = Alignment(horizontal="center")
        
        # Fecha Vencimiento
        c_fec_cell = ws_clients.cell(row=row_idx, column=8, value=c_fec)
        c_fec_cell.alignment = Alignment(horizontal="center")
        
        # Fórmula de Estado Vencimiento
        formula = f'=IF(H{row_idx}="","",IF(DATEVALUE(TEXT(H{row_idx},"yyyy-mm-dd"))<TODAY(),"VENCIDO",IF(DATEVALUE(TEXT(H{row_idx},"yyyy-mm-dd"))=TODAY(),"VENCE HOY",IF(DATEVALUE(TEXT(H{row_idx},"yyyy-mm-dd"))-TODAY()<=3,"URGENTE","EN TIEMPO"))))'
        c_formula_cell = ws_clients.cell(row=row_idx, column=9, value=formula)
        c_formula_cell.alignment = Alignment(horizontal="center")
        c_formula_cell.font = bold_font
        
        # Ultima_Fecha_Envio (inicialmente vacía)
        ws_clients.cell(row=row_idx, column=10, value="").alignment = Alignment(horizontal="center")
        
        # Estatus_Envio
        ws_clients.cell(row=row_idx, column=11, value=c_est).alignment = Alignment(horizontal="center")
        
        # Notas
        ws_clients.cell(row=row_idx, column=12, value=c_not)
        
        for c in range(1, 13):
            ws_clients.cell(row=row_idx, column=c).border = thin_border
            ws_clients.cell(row=row_idx, column=c).font = normal_font
    
    # Validaciones para Obligación
    dv_ob = DataValidation(
        type="list",
        formula1='"Pago Provisional ISR,Declaración Mensual IVA,Declaración Informativa DIOT,Retenciones de ISR / IVA,Declaración Anual,Nómina y Cuotas IMSS,Pago Estatal sobre Nómina,Otro"',
        allow_blank=True
    )
    ws_clients.add_data_validation(dv_ob)
    dv_ob.add("F5:F1000")
    
    # Ajuste automático de ancho de columnas en ambas hojas
    for ws in [ws_config, ws_clients]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if val.startswith("="):
                    val = "EN TIEMPO"
                if len(val) > max_len and cell.row > 2:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
    # Fijar panel superior en Obligaciones_Clientes
    ws_clients.freeze_panes = "A5"
    
    filename = "Plantilla_Recordatorio_Fiscales.xlsx"
    wb.save(filename)
    print(f"Plantilla generada exitosamente: {filename}")

if __name__ == "__main__":
    create_fiscal_reminder_template()
